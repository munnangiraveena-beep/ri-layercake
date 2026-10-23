"""PPTX structure: names, groups, scheme colours (no RGB), table cells, alt text."""

import io
import zipfile

from pptx import Presentation

import ri_layercake as lc
from ri_layercake import gallery


def export(tower, **kw):
    return tower.to_pptx(deterministic=True, **kw)


def slide_xml(data, n=1):
    return zipfile.ZipFile(io.BytesIO(data)).read(f"ppt/slides/slide{n}.xml").decode()


def test_layer_groups_names_and_alt_text():
    t = lc.Tower.from_text("10 ret, 40 xs 10, 50 xs 50 @ 75%", title="Renewal")
    prs = Presentation(io.BytesIO(export(t)))
    shapes = {s.name: s for s in prs.slides[0].shapes}
    assert {"Title", "Retention · 10 ret", "Layer 1 · 40 xs 10", "Layer 2 · 50 xs 50", "Layer table"} <= set(shapes)
    grp = shapes["Layer 2 · 50 xs 50"]
    assert grp.shape_type == 6  # group
    assert {s.name for s in grp.shapes} == {"L2.block", "L2.unplaced", "L2.label"}
    assert grp._element.nvGrpSpPr.cNvPr.get("descr") == "Layer 2: 50 million excess of 50 million, 75% placed"
    assert prs.slides[0].shapes.title.text == "Renewal"


def test_scheme_colours_only():
    for name in gallery.names():
        xml = slide_xml(export(gallery.load(name), scale="linear", loss=100))
        assert "srgbClr" not in xml
        assert "schemeClr" in xml
        has_unplaced = any(l.share < 1 for l in gallery.load(name).layers if l.kind != "retention")
        assert ("<a:pattFill" in xml) == has_unplaced


def test_native_table_with_totals():
    t = lc.Tower.from_text("10 ret, 40 xs 10 prem 4, 50 xs 50 prem 3")
    prs = Presentation(io.BytesIO(export(t)))
    table = next(s for s in prs.slides[0].shapes if s.has_table).table
    cells = [[c.text for c in row.cells] for row in table.rows]
    assert cells[0][0] == "Layer" and cells[1][0] == "Layer 2"
    assert cells[-1][:2] == ["Total", "90"]
    assert table.last_row is True


def test_theme_palettes_written_to_theme_xml():
    data = export(gallery.load("property_cat"), palette="ember")
    z = zipfile.ZipFile(io.BytesIO(data))
    theme = next(z.read(n).decode() for n in z.namelist() if n.startswith("ppt/theme/"))
    assert 'name="ri-layercake ember"' in theme and "9C3D1C" in theme


def test_deterministic_exports_are_identical():
    t = gallery.load("partly_placed")
    assert export(t) == export(t)


def test_engine_version_in_custom_properties():
    z = zipfile.ZipFile(io.BytesIO(export(gallery.load("casualty_xl"))))
    custom = z.read("docProps/custom.xml").decode()
    assert f"<vt:lpwstr>{lc.__version__}</vt:lpwstr>" in custom


def test_notes_carry_disclaimer():
    prs = Presentation(io.BytesIO(export(gallery.load("casualty_xl"))))
    assert "not pricing or actuarial advice" in prs.slides[0].notes_slide.notes_text_frame.text


def test_formula_injection_escaped():
    t = lc.Tower.from_text("40 xs 10", title="=HYPERLINK(1)")
    prs = Presentation(io.BytesIO(export(t)))
    assert prs.slides[0].shapes.title.text == "'=HYPERLINK(1)"


def test_template_pptx_and_potx(tmp_path):
    base = Presentation()
    base.slide_width, base.slide_height = 12_192_000, 6_858_000
    path = tmp_path / "house.pptx"
    base.save(path)
    data = gallery.load("property_cat").to_pptx(template=str(path), deterministic=True)
    assert Presentation(io.BytesIO(data)).slides[0].shapes.title.text == "Property cat XL 2027"

    potx = tmp_path / "house.potx"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(potx, "w") as dst:
        for item in src.infolist():
            blob = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                blob = blob.replace(b"presentation.main+xml", b"template.main+xml")
            dst.writestr(item, blob)
    data = gallery.load("property_cat").to_pptx(template=str(potx))
    assert len(Presentation(io.BytesIO(data)).slides) == 1


def test_table_slide_option():
    data = gallery.load("property_cat").to_pptx(table_slide=True)
    assert len(Presentation(io.BytesIO(data)).slides) == 2


def test_frames():
    for frame, size in [("square", (6_858_000, 6_858_000)), ("tall", (6_858_000, 12_192_000))]:
        prs = Presentation(io.BytesIO(gallery.load("casualty_xl").to_pptx(frame=frame)))
        assert (prs.slide_width, prs.slide_height) == size
