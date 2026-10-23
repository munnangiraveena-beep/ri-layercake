"""PPTX renderer: native shapes and a native table in theme colours.

Fills and lines use scheme colours with tint, never RGB, so pasting the
slide into a client deck recolours it to that deck's theme.
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from pathlib import Path
from typing import TYPE_CHECKING, Any, BinaryIO

from lxml import etree
from pptx import Presentation
from pptx.dml.color import ColorFormat
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_PATTERN, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.util import Emu, Pt

from . import fmt
from .scene import View
from .themes import FONTS, PALETTES

if TYPE_CHECKING:
    from pptx.presentation import Presentation as PresentationT

    from .model import Tower
    from .scene import Scene

DISCLAIMER = ("Illustrative only. Based on the inputs provided; not pricing or actuarial advice. "
              "Made with ri-layercake.")
FRAME_EMU = {"slide_16x9": (12_192_000, 6_858_000), "square": (6_858_000, 6_858_000),
             "tall": (6_858_000, 12_192_000)}
THEME = {
    "accent1": MSO_THEME_COLOR.ACCENT_1, "accent2": MSO_THEME_COLOR.ACCENT_2,
    "accent3": MSO_THEME_COLOR.ACCENT_3, "accent4": MSO_THEME_COLOR.ACCENT_4,
    "accent5": MSO_THEME_COLOR.ACCENT_5, "accent6": MSO_THEME_COLOR.ACCENT_6,
    "tx1": MSO_THEME_COLOR.TEXT_1, "bg1": MSO_THEME_COLOR.BACKGROUND_1,
    "tx2": MSO_THEME_COLOR.TEXT_2, "bg2": MSO_THEME_COLOR.BACKGROUND_2,
}
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PT_PER_PX = 0.75
FIXED_TIME = dt.datetime(2024, 1, 1, 0, 0, 0)
TEMPLATE_CT = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
PRESENTATION_CT = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"


def to_pptx(
    tower: Tower,
    path_or_buffer: str | Path | BinaryIO | None = None,
    view: View | None = None,
    template: str | Path | BinaryIO | bytes | None = None,
    layout: str | None = None,
    table_slide: bool = False,
    deterministic: bool = False,
    **options: Any,
) -> bytes | None:
    """Write a tower slide. Returns the bytes when no path or buffer is given."""
    from . import __version__

    view = view or View(**options)
    prs = _open(template, view)
    add_tower_slides(prs, tower, view, template is not None, layout, table_slide)
    _core_props(prs, tower, deterministic)
    _custom_props(prs, {"ri-layercake engine": __version__, "ri-layercake tower": tower.to_json()})
    return _save(prs, path_or_buffer)


def _save(prs: PresentationT, path_or_buffer: Any) -> bytes | None:
    if path_or_buffer is None:
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()
    prs.save(path_or_buffer)
    return None


def add_tower_slides(prs: PresentationT, tower: Tower, view: View, own_template: bool, layout: str | None = None,
                     table_slide: bool = False) -> Any:
    """Add the tower slide (and optionally a full-size table slide) to a presentation."""
    scene = tower.scene(view)
    slide_layout = _pick_layout(prs, layout)
    slide = prs.slides.add_slide(slide_layout)
    w_emu, h_emu = prs.slide_width, prs.slide_height
    k = min(w_emu / scene["frame"]["w"], h_emu / scene["frame"]["h"])
    ox = (w_emu - scene["frame"]["w"] * k) / 2
    oy = (h_emu - scene["frame"]["h"] * k) / 2
    r = _Renderer(slide, tower, k, ox, oy)

    title_ph = _keep_title(slide, bool(tower.title))
    if title_ph is not None and not own_template:
        # Default template: put the title placeholder where the scene draws it.
        title_ph.left, title_ph.top = r.X(40), r.Y(36)
        title_ph.width, title_ph.height = r.L(scene["frame"]["w"] - 80), r.L(40)
        tf = title_ph.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
    r.draw(scene, title_ph)
    if title_ph is not None and not own_template:
        _style_title(title_ph, r)

    if table_slide and scene["table"]["rows"]:
        s2 = prs.slides.add_slide(slide_layout)
        t2 = _keep_title(s2, True)
        if t2 is not None:
            t2.text = fmt.escape(f"{tower.title or 'Programme'}: layers")
        r2 = _Renderer(s2, tower, k, ox, oy)
        W = scene["frame"]["w"]
        rh = 28.0
        box = {"x": 40, "y": 120, "w": W - 80, "h": rh * (len(scene["table"]["rows"]) + 2), "row_h": rh}
        from .scene import _col_widths
        from .table import table_data

        table = dict(scene["table"], box=box, col_w=_col_widths(table_data(tower), box["w"]))
        r2.table(table)
        _notes(s2)
    _notes(slide)
    return slide


def _keep_title(slide: Any, has_title: bool) -> Any:
    """Remove every placeholder except the title (and that too when there is no title)."""
    title_ph = slide.shapes.title
    for ph in list(slide.placeholders):
        if title_ph is None or ph.shape_id != title_ph.shape_id or not has_title:
            ph._element.getparent().remove(ph._element)
    return title_ph if has_title else None


def _style_title(title_ph: Any, r: Any) -> None:
    p = title_ph.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for run in p.runs:
        run.font.size = r.pt(28)
        run.font.bold = True


def add_table_slide(prs: PresentationT, title: str, columns: list[str], rows: list[list[str]],
                    totals: list[str] | None = None, layout: str | None = None, own_template: bool = False) -> Any:
    """A slide with a title and one native table, e.g. the portfolio summary of a book."""
    from .model import Tower

    slide = prs.slides.add_slide(_pick_layout(prs, layout))
    title_ph = _keep_title(slide, True)
    k = prs.slide_width / 1280
    r = _Renderer(slide, Tower(), k, 0, (prs.slide_height - 720 * k) / 2)
    if title_ph is not None:
        title_ph.text = fmt.escape(title)
        if not own_template:
            title_ph.left, title_ph.top = r.X(40), r.Y(36)
            title_ph.width, title_ph.height = r.L(1200), r.L(40)
            _style_title(title_ph, r)
    rh = 30.0
    n = len(rows) + 1 + (1 if totals else 0)
    weights = [max(len(c), *(len(row[j]) for row in rows)) + 4 for j, c in enumerate(columns)]
    widths = [1200 * w / sum(weights) for w in weights]
    table = {"columns": columns, "rows": rows, "totals": totals or [""] * len(columns),
             "box": {"x": 40, "y": 130, "w": 1200, "h": rh * n, "row_h": rh, "size": 14}, "col_w": widths}
    if not totals:
        table["rows"] = rows
    r.table(table)
    _notes(slide)
    return slide


# ---------------------------------------------------------------------- setup

def _open(template: Any, view: View) -> PresentationT:
    if template is None:
        prs = Presentation()
        prs.slide_width, prs.slide_height = (Emu(v) for v in FRAME_EMU[view.frame])
        _apply_palette(prs, view.palette)
        return prs
    data = template if isinstance(template, bytes) else (
        template.read() if hasattr(template, "read") else Path(template).read_bytes())
    return Presentation(io.BytesIO(_potx_to_pptx(data)))


def _potx_to_pptx(data: bytes) -> bytes:
    """python-pptx opens .pptx only; a .potx differs just in its main content type."""
    src = zipfile.ZipFile(io.BytesIO(data))
    ct = src.read("[Content_Types].xml")
    if TEMPLATE_CT.encode() not in ct:
        return data
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            blob = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                blob = blob.replace(TEMPLATE_CT.encode(), PRESENTATION_CT.encode())
            dst.writestr(item, blob)
    return out.getvalue()


def _apply_palette(prs: PresentationT, palette: str) -> None:
    """Write a built-in palette into the default template's theme XML."""
    pal = PALETTES.get(palette, PALETTES["ocean"])
    theme_part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    scheme = root.find(f".//{{{A}}}clrScheme")
    assert scheme is not None
    scheme.set("name", f"ri-layercake {palette}")
    for slot, hexv in pal.items():
        el = scheme.find(f"{{{A}}}{slot}")
        if el is None:
            continue
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{{{A}}}srgbClr", val=hexv)
    fonts = root.find(f".//{{{A}}}fontScheme")
    if fonts is not None:
        for which, tag in (("major", "majorFont"), ("minor", "minorFont")):
            latin = fonts.find(f"{{{A}}}{tag}/{{{A}}}latin")
            if latin is not None:
                latin.set("typeface", FONTS[which])
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _pick_layout(prs: PresentationT, name: str | None) -> Any:
    layouts = list(prs.slide_layouts)
    if name:
        for l in layouts:
            if l.name.lower() == name.lower():
                return l
        raise ValueError(f"No layout named {name!r}; found {[l.name for l in layouts]}")
    for l in layouts:
        if l.name.lower() == "title only":
            return l
    for l in layouts:
        types = [ph.placeholder_format.type for ph in l.placeholders]  # type: ignore[misc]
        if len(types) <= 4 and any("TITLE" in str(t) for t in types):
            return l
    return layouts[-1]


def _notes(slide: Any) -> None:
    slide.notes_slide.notes_text_frame.text = DISCLAIMER


def _core_props(prs: PresentationT, tower: Tower, deterministic: bool) -> None:
    cp = prs.core_properties
    now = FIXED_TIME if deterministic else dt.datetime.now(dt.timezone.utc).replace(tzinfo=None, microsecond=0)
    cp.created = cp.modified = now
    cp.last_printed = now
    cp.revision = 1
    cp.author = cp.last_modified_by = "ri-layercake"
    cp.title = tower.title or "Reinsurance tower"


def _custom_props(prs: PresentationT, props: dict[str, str]) -> None:
    from xml.sax.saxutils import escape

    body = "".join(
        f'<property fmtid="{{D5CDD505-2E9C-101B-9397-08002B2CF9AE}}" pid="{i}" name="{escape(k)}">'
        f"<vt:lpwstr>{escape(v)}</vt:lpwstr></property>"
        for i, (k, v) in enumerate(props.items(), start=2)
    )
    xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           f"{body}</Properties>").encode()
    package = prs.part.package
    try:
        package.part_related_by(RT.CUSTOM_PROPERTIES)._blob = xml
        return
    except KeyError:
        pass
    part = Part(PackURI("/docProps/custom.xml"),
                "application/vnd.openxmlformats-officedocument.custom-properties+xml", package, xml)
    package.relate_to(part, RT.CUSTOM_PROPERTIES)


# ---------------------------------------------------------------------- drawing

class _Renderer:
    def __init__(self, slide: Any, tower: Tower, k: float, ox: float, oy: float):
        self.slide, self.tower, self.k, self.ox, self.oy = slide, tower, k, ox, oy

    def X(self, px: float) -> Emu:
        return Emu(round(self.ox + px * self.k))

    def Y(self, px: float) -> Emu:
        return Emu(round(self.oy + px * self.k))

    def L(self, px: float) -> Emu:
        return Emu(max(1, round(px * self.k)))

    def pt(self, px: float) -> Pt:
        return Pt(round(px * PT_PER_PX * self.k / 9525, 1))

    def draw(self, scene: Scene, title_ph: Any) -> None:
        groups: dict[str, Any] = {}
        order: list[str] = []
        for s in scene.shapes:
            if s.get("export") is False:
                continue
            key = s.get("layer") or s["role"]
            if s["role"] in ("title", "subtitle") and title_ph is not None:
                if s["role"] == "title":
                    title_ph.text = fmt.escape(s["runs"][0]["t"])
                    title_ph.name = "Title"
                    continue
            if key not in groups:
                groups[key] = []
                order.append(key)
            groups[key].append(s)

        for key in order:
            shapes = groups[key]
            if key in ("title", "subtitle"):
                for s in shapes:
                    self.shape(self.slide.shapes, s).name = key.capitalize()
                continue
            name, alt = self._group_name(key, shapes)
            grp = self.slide.shapes.add_group_shape()
            for s in shapes:
                self.shape(grp.shapes, s)
            grp.name = name
            if alt:
                grp._element.nvGrpSpPr.cNvPr.set("descr", alt)

        if scene.data.get("show_table"):
            self.table(scene["table"])

    def _group_name(self, key: str, shapes: list[dict[str, Any]]) -> tuple[str, str | None]:
        tower = self.tower
        try:
            layer = tower.layer(key)
        except KeyError:
            names = {"axis": "Axis", "grid": "Gridlines", "axis_label": "Axis labels", "loss": "Loss line",
                     "loss_callout": "Loss callout", "gap": "Gap", "gap_label": "Gap label", "empty": "Empty",
                     "caption": "Captions", "leader": "Leader lines"}
            return names.get(key, key.replace("_", " ").capitalize()), None
        short = fmt.short_label(tower, layer)
        if layer.kind == "retention":
            return (f"Retention · {short}",
                    f"Retention: the cedant keeps the first {fmt.words(tower, layer.limit)}")
        if layer.kind == "qs":
            return (f"Quota share · {fmt.pct(layer.cession)}",
                    f"Quota share: {fmt.pct(layer.cession)} of each risk ceded, up to a limit of "
                    f"{fmt.words(tower, layer.limit)}; {fmt.pct(layer.share)} placed")
        if layer.kind == "sl":
            return (f"{tower.display_name(layer)} · {short}",
                    f"{tower.display_name(layer)}: {fmt.pct(layer.limit)} excess of {fmt.pct(layer.attachment)} "
                    f"of EPI, annual aggregate; {fmt.pct(layer.share)} placed")
        n = tower.layer_number(layer.id)
        alt = (f"Layer {n}: {fmt.words(tower, layer.limit)} excess of {fmt.words(tower, layer.attachment)}, "
               f"{fmt.pct(layer.share)} placed")
        if layer.name:
            alt = f"{fmt.escape(layer.name)}. " + alt
        return f"Layer {n} · {short}", alt

    def shape(self, coll: Any, s: dict[str, Any]) -> Any:
        t = s["type"]
        if t == "rect":
            shp = coll.add_shape(MSO_SHAPE.RECTANGLE, self.X(s["x"]), self.Y(s["y"]), self.L(s["w"]), self.L(s["h"]))
            shp.name = s["id"]
            shp.shadow.inherit = False
            fill = s.get("fill")
            if not fill:
                shp.fill.background()
            elif fill.get("pattern"):
                shp.fill.patterned()
                shp.fill.pattern = MSO_PATTERN.WIDE_UPWARD_DIAGONAL
                _theme(shp.fill.fore_color, fill["theme"], fill.get("tint", 0))
                _theme(shp.fill.back_color, "bg1", 0)
            else:
                shp.fill.solid()
                _theme(shp.fill.fore_color, fill["theme"], fill.get("tint", 0))
            self._line(shp.line, s.get("stroke"))
            if shp.has_text_frame:
                shp.text_frame.text = ""
            return shp
        if t == "line":
            shp = coll.add_connector(MSO_CONNECTOR.STRAIGHT, self.X(s["x1"]), self.Y(s["y1"]),
                                     self.X(s["x2"]), self.Y(s["y2"]))
            shp.name = s["id"]
            self._line(shp.line, s.get("stroke") or {"theme": "tx1", "w": 1})
            return shp
        if t == "text":
            return self.text(coll, s)
        raise ValueError(f"Unknown shape type {t}")

    def _line(self, line: Any, stroke: dict[str, Any] | None) -> None:
        if not stroke or stroke.get("status"):
            line.fill.background()
            return
        line.fill.solid()
        _theme(line.color, stroke["theme"], stroke.get("tint", 0))
        line.width = self.pt(stroke.get("w", 1))
        if stroke.get("dash"):
            line.dash_style = MSO_LINE_DASH_STYLE.DASH

    def text(self, coll: Any, s: dict[str, Any]) -> Any:
        runs = s["runs"]
        sizes = [r.get("s", s.get("size", 14)) for r in runs]
        w = max(fmt_len(r["t"]) * sz * 0.56 for r, sz in zip(runs, sizes)) + 12
        h = sum(sz * 1.25 for sz in sizes) + 4
        anchor = s.get("anchor", "start")
        left = s["x"] - (w / 2 if anchor == "middle" else w if anchor == "end" else 0)
        top = s["y"] - sizes[0]
        box = coll.add_textbox(self.X(left), self.Y(top), self.L(w), self.L(h))
        box.name = s["id"]
        tf = box.text_frame
        tf.word_wrap = False
        tf.auto_size = None
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        align = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}[anchor]
        colour = s.get("fill") or {"theme": "tx1"}
        major = s.get("role") == "title"
        for i, (run, size) in enumerate(zip(runs, sizes)):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = fmt.escape(run["t"])
            f = r.font
            f.size = self.pt(size)
            f.bold = bool(run.get("b") or run.get("warn"))
            f.name = "+mj-lt" if major else "+mn-lt"
            _theme(f.color, colour["theme"], colour.get("tint", 0))
        return box

    def table(self, table: dict[str, Any]) -> Any:
        box = table["box"]
        cols = table["columns"]
        rows = [cols] + table["rows"] + [table["totals"]]
        gf = self.slide.shapes.add_table(len(rows), len(cols), self.X(box["x"]), self.Y(box["y"]),
                                         self.L(box["w"]), self.L(box["h"]))
        gf.name = "Layer table"
        tbl = gf.table
        tbl.first_row = True
        tbl.last_row = True
        tbl.horz_banding = True
        for j, w in enumerate(table["col_w"]):
            tbl.columns[j].width = self.L(w)
        size = self.pt(box.get("size", min(13.0, box["row_h"] * 0.55)))
        for i, row in enumerate(rows):
            tbl.rows[i].height = self.L(box["row_h"])
            for j, value in enumerate(row):
                cell = tbl.cell(i, j)
                cell.margin_top = cell.margin_bottom = Emu(0)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.text = fmt.escape(value)
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
                for r in p.runs:
                    r.font.size = size
                    r.font.name = "+mn-lt"
                    if i == len(rows) - 1:
                        r.font.bold = True
        return gf


def fmt_len(text: str) -> int:
    return len(text)


def _theme(color: ColorFormat, theme: str, tint: float) -> None:
    color.theme_color = THEME[theme]
    if tint:
        color.brightness = round(float(tint), 2)
